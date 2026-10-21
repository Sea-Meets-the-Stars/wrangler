""" Build claude_prompts/llc4320_v2/llc4320_v2_design.pptx: the design for adding
the remaining LLC4320 v2 surface fields (summary of llc4320_v2_design.md).

    python claude_prompts/llc4320_v2/make_llc4320_v2_design_slides.py

Needs python-pptx.  Every slide carries a one-line takeaway under its title.
Every font is at least MIN_PT (20 pt); `_run` raises if a smaller size slips
in.  Re-running overwrites hand edits made in PowerPoint/LibreOffice.
"""

import os

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, 'llc4320_v2_design.pptx')

# Palette: deep-ocean navy dominates, warm coral as the single accent (same as
# docs/slides/llc4320_v2.pptx so the two decks read as one series).
NAVY = RGBColor(0x0B, 0x2A, 0x3F)
TEAL = RGBColor(0x1F, 0x6F, 0x9F)
ICE = RGBColor(0xDC, 0xEB, 0xF2)
PALE = RGBColor(0xEE, 0xF3, 0xF7)
CORAL = RGBColor(0xE4, 0x57, 0x2E)
BLUSH = RGBColor(0xFB, 0xE4, 0xDC)
INK = RGBColor(0x1C, 0x26, 0x30)
MUTED = RGBColor(0x5B, 0x6B, 0x78)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GRID = RGBColor(0xDD, 0xE3, 0xE8)
HEAD, BODY, MONO = 'Cambria', 'Calibri', 'Courier New'

MIN_PT = 20
BODY_PT = 20
TITLE_PT = 34
W, H = 13.333, 7.5
DATE = '2026-10-20'


# ------------------------------------------------------------------ helpers

def _run(p, text, size, color=INK, bold=False, italic=False, font=BODY):
    if size < MIN_PT:
        raise ValueError(f'{size} pt is below the {MIN_PT} pt minimum: {text!r}')
    r = p.add_run()
    r.text = text
    f = r.font
    f.size, f.bold, f.italic, f.name = Pt(size), bold, italic, font
    f.color.rgb = color
    return r


def text(slide, x, y, w, h, paras, size=BODY_PT, color=INK, font=BODY, anchor=MSO_ANCHOR.TOP,
         align=PP_ALIGN.LEFT, space_after=8, bullets=False, name=None):
    """Text box; *paras* is a list of str or list of (text, {opts}) runs."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for n, para in enumerate(paras):
        p = tf.paragraphs[0] if n == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        runs = [(para, {})] if isinstance(para, str) else para
        if bullets:
            runs = [('•  ', {'color': CORAL, 'bold': True})] + list(runs)
        for t, o in runs:
            _run(p, t, o.get('size', size), o.get('color', color), o.get('bold', False),
                 o.get('italic', False), o.get('font', font))
    return tb


def box(slide, x, y, w, h, fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE, name=None):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        s.name = name
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    s.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = 0.06
    return s


def dot(slide, x, y, label, fill, d=0.5):
    o = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    o.fill.solid()
    o.fill.fore_color.rgb = fill
    o.line.fill.background()
    tf = o.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    _run(p, str(label), 20, WHITE, bold=True)
    return o


def arrow(slide, x, y, w, h=0.35):
    a = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(h))
    a.fill.solid()
    a.fill.fore_color.rgb = CORAL
    a.line.fill.background()
    return a


def background(slide, color):
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = color


def title(slide, t, summary, dark=False, size=TITLE_PT):
    """Title plus the one-line takeaway every slide carries."""
    text(slide, 0.6, 0.35, W - 1.2, 0.75, [t], size=size, color=WHITE if dark else NAVY,
         font=HEAD, name='title')
    text(slide, 0.6, 1.1, W - 1.2, 0.5, [summary], size=20, color=ICE if dark else CORAL,
         name='summary')


def footer(slide, n, dark=False):
    text(slide, W - 1.4, H - 0.55, 0.9, 0.35, [str(n)], size=20, color=ICE if dark else MUTED,
         align=PP_ALIGN.RIGHT, name='slide number')


def stat(slide, x, y, w, value, label, color=NAVY, label_color=MUTED, size=36):
    text(slide, x, y, w, 0.7, [value], size=size, color=color, font=HEAD, space_after=0)
    text(slide, x, y + 0.7, w, 0.5, [label], size=20, color=label_color, space_after=0)


def notes(slide, t):
    slide.notes_slide.notes_text_frame.text = t


def card(slide, x, y, w, h, head, items, fill=ICE, head_color=NAVY, body_color=INK,
         head_size=24, body_size=20, space_after=10):
    box(slide, x, y, w, h, fill)
    text(slide, x + 0.3, y + 0.22, w - 0.6, 0.55, [head], size=head_size, color=head_color,
         font=HEAD)
    text(slide, x + 0.3, y + 0.9, w - 0.6, h - 1.1, items, size=body_size, color=body_color,
         bullets=True, space_after=space_after)


def table(slide, x, y, w, rows, col_w, row_h=0.5, header_fill=NAVY, size=20):
    """Simple table; *rows* is a list of lists of str; first row is the header."""
    shp = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w),
                                 Inches(row_h * len(rows)))
    tbl = shp.table
    for j, cw in enumerate(col_w):
        tbl.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        tbl.rows[i].height = Inches(row_h)
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.margin_left = cell.margin_right = Inches(0.1)
            cell.margin_top = cell.margin_bottom = Inches(0.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = header_fill if i == 0 else (WHITE if i % 2 else PALE)
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT
            _run(p, val, size, WHITE if i == 0 else INK, bold=(i == 0))
    return shp


def new_slide(prs, dark=False):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    background(s, NAVY if dark else WHITE)
    return s


# ------------------------------------------------------------------ slides

def s_title(prs):
    s = new_slide(prs, dark=True)
    box(s, 7.6, 1.3, 5.1, 4.9, TEAL, name='field grid card')
    names = ['Theta  Salt  U  V', 'Eta  KPPhbl  PhiBot', 'oceQnet  oceQsw', 'oceFWflx  oceSflux',
             'oceTAUX  oceTAUY', 'SHIfwFlx  SHIhtFlx', 'SIarea  SIheff  SIhsnow',
             'SIhsalt  SIuice  SIvice']
    text(s, 7.95, 1.55, 4.5, 4.5, names, size=22, color=WHITE, font=MONO, space_after=5)
    text(s, 0.6, 1.5, 6.6, 2.0, ['LLC4320 v2:', 'the other surface fields'], size=40,
         color=WHITE, font=HEAD, space_after=0)
    text(s, 0.6, 3.7, 6.6, 1.6,
         ['Design for adding 20 hourly 2D fields to s3://llc4320-v2/SURFACE '
          'without rewriting the SST already there'], size=22, color=ICE)
    text(s, 0.6, 5.6, 6.6, 0.6, [f'J. X. Prochaska  ·  {DATE}  ·  design, not yet implemented'],
         size=20, color=ICE)
    footer(s, 1, dark=True)
    notes(s, 'Summary of claude_prompts/llc4320_v2/llc4320_v2_design.md. Sections 1-3 of that '
             'document describe the existing SST product; 4-11 the design for the new fields.')


def s_goal(prs, n):
    s = new_slide(prs)
    title(s, 'Goal and where we stand',
          'SST is published; twenty more fields go into the same hourly stores, one pass, in parallel')
    text(s, 0.6, 1.85, 6.3, 4.9, [
        [('Published: ', {'bold': True}), ('hourly global SST, 2023-01-01 to 2024-01-31, public', {})],
        [('Missing: ', {'bold': True}), ('2 outage hours + 189 hours in a closed folder', {})],
        [('Add: ', {'bold': True}), ('SSS, SSU, SSV, Eta, fluxes, ice shelf, sea ice', {})],
        [('Rule 1: ', {'bold': True}), ('never rewrite the Theta already on the bucket', {})],
        [('Rule 2: ', {'bold': True}), ('one writer per store, enforced by the code', {})],
        [('Decide first: ', {'bold': True, 'color': CORAL}),
         ('storage quota, record length', {'color': CORAL})],
    ], size=22, bullets=True, space_after=14)
    box(s, 7.3, 1.8, 5.4, 5.0, ICE, name='status card')
    stat(s, 7.7, 2.0, 2.4, '9,501', 'hourly SST stores')
    stat(s, 10.2, 2.0, 2.4, '4.4 TB', 'on the bucket')
    stat(s, 7.7, 3.55, 2.4, '20', 'new fields', color=CORAL)
    stat(s, 10.2, 3.55, 2.4, '95', 'raw folders (was 83)', color=CORAL)
    stat(s, 7.7, 5.1, 2.4, '1 pass', 'all fields together')
    stat(s, 10.2, 5.1, 2.4, 'N procs', 'disjoint date windows')
    footer(s, n)
    notes(s, 'Decisions from the Q&A: preserve existing data (mode=a), all fields present in every '
             'folder, one pass, native SSU/SSV only, full record, one store per hour, hFac masks, '
             'verify configurable, benchmark before launch. The 95-folder count and the storage '
             'total are the two things that still need a decision.')


def s_source(prs, n):
    s = new_slide(prs)
    title(s, 'The raw source on Pleiades',
          'Every surface field is a plain compact binary, one file per hour, dated by position in its folder')
    card(s, 0.6, 1.85, 4.0, 4.9, 'Folders', [
        'dbwhitt/llc_4320/OUT on nobackupp27',
        'Named by the segment they cover: YYYY_MM_DD_HHMMSS_to_…',
        '95 folders today; 83 when SST ran',
        'Last one: symlink into the live run',
    ])
    card(s, 4.8, 1.85, 4.0, 4.9, 'Files', [
        '<FIELD>.<iteration>.data',
        '4320 × 56160 big-endian float32',
        '970,444,800 bytes, no header, no .meta',
        '21 prefixes in all 95 folders',
    ])
    card(s, 9.0, 1.85, 3.75, 4.9, 'Dating', [
        'date = folder start + (n+1) h',
        'files per folder = span in hours',
        'ΔT 5–20 s varies; iterations only cross-check',
        '9,692 h proven gapless (so far)',
    ])
    footer(s, n)
    notes(s, 'Dan Whitt, 2026-09-09: hourly snapshots; first file one hour after the folder start; '
             'last at the folder end; no hour in two folders. Masks: Bron Nelson\'s '
             'bitmask.4320x173/hFacC,S,W.bits, level 0. Grid: grid_90x90x19493. The 3D .shrunk '
             'output and the R_*, *F, Eta_tide* families are out of scope.')


def s_pipeline(prs, n):
    s = new_slide(prs)
    title(s, 'How SST was extracted (and how every field will be)',
          'Read, reorder faces, mask, compress, write one Zarr store per hour, verify, consolidate')
    steps = [
        ('Discover', 'folders → dated hours', True),
        ('Read', '.data → (13, 4320, 4320)', False),
        ('Faces', 'compact_to_faces', False),
        ('Mask', 'hFac bits → NaN land', False),
        ('Write', 'zstd chunks (1,720,720)', False),
        ('Verify', 'read back, mark complete', True),
    ]
    bw, gap, y = 1.85, 0.25, 1.9
    for k, (head, body, dark) in enumerate(steps):
        x = 0.6 + k * (bw + gap)
        box(s, x, y, bw, 2.0, NAVY if dark else ICE)
        text(s, x + 0.15, y + 0.18, bw - 0.3, 0.5, [head], size=22, font=HEAD,
             color=WHITE if dark else NAVY)
        text(s, x + 0.15, y + 0.8, bw - 0.3, 1.1, [body], size=20, color=ICE if dark else INK)
        if k < len(steps) - 1:
            arrow(s, x + bw + 0.03, y + 0.85, gap - 0.06)
    text(s, 0.6, 4.3, 6.0, 2.5, [
        [('Store: ', {'bold': True}), ('YYYYMMDDTHH.zarr, float32, NaN on land', {})],
        [('Size: ', {'bold': True}), ('~470 MB and ~400 objects per field-hour', {})],
        [('Compression: ', {'bold': True}), ('zstd, 1.19× on ocean points (measured)', {})],
    ], size=20, bullets=True, space_after=10)
    text(s, 6.9, 4.3, 5.9, 2.5, [
        [('Retries: ', {'bold': True}), ('4 attempts, backoff; a failed store is left incomplete', {})],
        [('Idempotent: ', {'bold': True}), ('complete stores skipped on re-run', {})],
        [('Metadata: ', {'bold': True}), ('consolidated, so one request opens a store', {})],
    ], size=20, bullets=True, space_after=10)
    footer(s, n)
    notes(s, 'wr_llc_v2_surface $OUT s3://llc4320-v2/SURFACE --fields SST --mask-dir --grid-dir. '
             'Zarr v3 over an async s3fs FsspecStore, path-style addressing. All-NaN (land) chunks '
             'are never written. Group attrs: model, selected_iteration, selected_date_utc, '
             'source_folder, FS, variables, face_layout, complete.')


def s_faces(prs, n):
    s = new_slide(prs)
    title(s, 'The face-order bug, found and repaired in place',
          'Faces 7–12 were scrambled by a naive reshape; values were only permuted, and all stores are fixed')
    # Schematic of the compact layout: facets 1-2 tall, facet 3 square, facets 4-5 wide
    x0, y0, u = 0.7, 1.9, 0.95
    box(s, x0, y0, u, 3 * u, TEAL, shape=MSO_SHAPE.RECTANGLE, name='facet 1')
    box(s, x0 + u + 0.1, y0, u, 3 * u, TEAL, shape=MSO_SHAPE.RECTANGLE, name='facet 2')
    box(s, x0 + 2 * u + 0.2, y0 + 2 * u, u, u, NAVY, shape=MSO_SHAPE.RECTANGLE, name='facet 3')
    box(s, x0 + 3 * u + 0.3, y0 + 2 * u, 3 * u, u, CORAL, shape=MSO_SHAPE.RECTANGLE, name='facet 4')
    box(s, x0 + 3 * u + 0.3, y0 + 0.8 * u, 3 * u, u, CORAL, shape=MSO_SHAPE.RECTANGLE, name='facet 5')
    text(s, x0, y0 + 3 * u + 0.2, 6.2, 1.2, [
        [('Facets 4–5 (faces 7–12) are stored as (4320, 3×4320) blocks; ', {}),
         ('their true faces are the three column blocks', {'bold': True})]], size=20, color=MUTED)
    text(s, 7.1, 1.85, 5.7, 5.0, [
        [('Cause: ', {'bold': True}), ('reshape straight to (13, 4320, 4320)', {})],
        [('Hidden: ', {'bold': True}), ('masks were reshaped the same way, so NaN checks passed', {})],
        [('Reader fix: ', {'bold': True}), ('compact_to_faces; matches xmitgcm and v1\'s grid', {})],
        [('Repair: ', {'bold': True}), ('5,492 stores rewritten in place, layout inferred from data first', {})],
        [('Marker: ', {'bold': True}), ('face_layout = "llc_faces" on every store', {})],
    ], size=20, bullets=True, space_after=12)
    footer(s, n)
    notes(s, 'Found 2026-10-02 from a Gulf Stream zoom. Repaired 2026-10-09 with wr_llc_v2_fix_faces, '
             '6 workers, 7.48 h. Verified: NaN patterns match grid.zarr maskC at 100.000 percent; '
             'median row-to-row SST step on faces 7-12 fell from 27.5 to 0.003-0.009 degC. Same '
             'fix applies to every new field automatically.')


def s_grid(prs, n):
    s = new_slide(prs)
    title(s, 'Store layout and grid.zarr',
          'One store per hour holds every variable; grid.zarr carries the C-grid staggering xgcm needs')
    table(s, 0.6, 1.85, 6.4, [
        ['Item', 'Value'],
        ['Store name', 'YYYYMMDDTHH.zarr'],
        ['Dims', '(face, j, i) or staggered j_g / i_g'],
        ['Chunks', '(1, 720, 720), float32, NaN fill'],
        ['grid.zarr', '23 variables, 4.8 GB, COMODO attrs'],
        ['Metadata', 'consolidated after every write'],
    ], col_w=[1.9, 4.5], row_h=0.62)
    card(s, 7.3, 1.85, 5.45, 4.9, 'Arakawa C-grid placement', [
        'Tracers, Eta, fluxes, SIarea: cell centre (j, i)',
        'U, oceTAUX, SIuice: west face (j, i_g)',
        'V, oceTAUY, SIvice: south face (j_g, i)',
        'Corners (j_g, i_g): XG, YG, rAz',
        'Fixed 2026-10-20; xgcm.Grid builds from the public store',
    ], space_after=9)
    footer(s, n)
    notes(s, 'VAR_HDIMS in llc_v2.py maps variable to horizontal dims, matching v1. '
             '_ensure_coords adds j_g/i_g to an hourly store the first time a staggered variable '
             'lands in it. Both shifts are -0.5 (MITgcm west/south convention).')


def s_access(prs, n):
    s = new_slide(prs)
    title(s, 'Public access and the reader',
          'Anonymous read-only S3 plus wrangler helpers that add the missing time axis')
    card(s, 0.6, 1.85, 6.0, 4.9, 'Bucket', [
        'https://s3-west.nrp-nautilus.io',
        's3://llc4320-v2/SURFACE, no credentials',
        'GetObject + ListBucket only; no public write',
        'wr_llc_v2_check_access verifies the policy',
    ])
    card(s, 6.8, 1.85, 5.95, 4.9, 'wrangler.ogcm.llc_v2_access', [
        'open_grid, open_store, open_range(start, end)',
        'nearest_index, extract_point, region_slices',
        'HOWTO: docs/howto/llc4320_v2.rst (Read the Docs)',
        [('New: ', {'bold': True}), ('open_range tolerates mixed variable sets', {})],
    ])
    footer(s, n)
    notes(s, 'During the weeks the new fields are being written, stores will differ in variable '
             'sets. open_range without variables= will keep the common set with a warning instead '
             'of failing in xr.concat.')


def s_fields(prs, n):
    s = new_slide(prs)
    title(s, 'The 20 new fields, in three groups',
          'Twelve dense fields cost what SST did; eight sparse ice fields are cheap in bytes but not in requests')
    face = {'color': MUTED}
    card(s, 0.6, 1.85, 3.4, 4.9, 'State (snapshots)', [
        [('SSS → Salt', {'font': MONO})],
        [('SSU → U', {'font': MONO}), ('   west face', face)],
        [('SSV → V', {'font': MONO}), ('   south face', face)],
        [('Eta', {'font': MONO})],
    ], space_after=12)
    card(s, 4.15, 1.85, 4.2, 4.9, 'Dense diagnostics', [
        [('KPPhbl  PhiBot', {'font': MONO})],
        [('oceQnet  oceQsw', {'font': MONO})],
        [('oceFWflx  oceSflux', {'font': MONO})],
        [('oceTAUX  oceTAUY', {'font': MONO})],
        [('west and south faces', face)],
    ], space_after=10)
    card(s, 8.5, 1.85, 4.25, 4.9, 'Sparse (ice)', [
        [('SHIfwFlx  SHIhtFlx', {'font': MONO})],
        [('SIarea  SIheff', {'font': MONO})],
        [('SIhsnow  SIhsalt', {'font': MONO})],
        [('SIuice  SIvice', {'font': MONO})],
        [('west and south faces', face)],
    ], fill=BLUSH, head_color=CORAL, space_after=10)
    footer(s, n)
    notes(s, 'All 21 prefixes are present in all 95 folders (inventory 2026-10-20). Store names '
             'follow xmitgcm/dbof. SHIfwFlx, SHIhtFlx, SIhsalt are new to SURFACE_FIELDS. '
             'Sparse means zero away from sea ice or ice shelves; their compression is inferred, '
             'to be measured by benchmark run D. SSU/SSV are published unrotated; the HOWTO gets '
             'the rotation example with CS/SN.')


def s_inventory(prs, n):
    s = new_slide(prs)
    title(s, 'What the full inventory changed',
          'The tree has 95 folders, not 83, so the record is probably longer; other file families are out of scope')
    card(s, 0.6, 1.85, 6.0, 4.9, '95 folders', [
        'SST ran over 83 (plus the closed live-run folder)',
        'Probably ~11 new segments: the run went on past 2024-02-08',
        'Inference: ~1,300–1,500 more hours, ~11,000 total',
        'One paste settles it: folder span + dry-run count',
        'The dating proof (count = span) must be redone',
    ], fill=BLUSH, head_color=CORAL, space_after=9)
    card(s, 6.8, 1.85, 5.95, 4.9, 'Not in scope (partial coverage)', [
        'R_Eta, R_Theta, R_U … in 5 folders: a second stream',
        'EtaF, ThetaF, SaltF …: 1–4 folders, other cadence',
        'Eta_tide (57), Eta_tide_presal (29): tidal SSH',
        '*.shrunk, every5, sections5deg, bloat: the 3D phase',
        'Ignored by construction: not present in every folder',
    ], space_after=9)
    footer(s, n)
    notes(s, 'The paste omitted the inventory\'s first lines (folder count, first and last date, '
             'unreadable folders), so the extra folders cannot be identified from here. What the '
             'other families are is inference from their names. The field resolution only keeps '
             'prefixes present in every readable folder, so nothing needs excluding by name.')


def s_sampling(prs, n):
    s = new_slide(prs)
    title(s, 'Open question: snapshots or hourly means?',
          'The diagnostics may be averaged over the hour; the design records it and can fix attributes later')
    text(s, 0.6, 1.85, 6.2, 4.9, [
        [('Certain: ', {'bold': True}), ('SST, SSS, SSU, SSV, Eta are snapshots', {})],
        [('Unknown: ', {'bold': True}), ('pkg/diagnostics fields average unless frequency < 0', {})],
        [('Checked: ', {'bold': True}), ('no .meta files, so no timeInterval to read', {})],
        [('Next: ', {'bold': True}), ('grep frequency in data.diagnostics / STDOUT.0000', {})],
        [('Or: ', {'bold': True}), ('ask Dan', {})],
    ], size=22, bullets=True, space_after=14)
    card(s, 7.1, 1.85, 5.65, 4.9, 'How the design handles it', [
        '--time-mean lists the averaged prefixes',
        'cell_methods = "time: mean (interval: 1 hour)"',
        'Until known: attribute says "not verified"',
        'Afterwards: metadata-only fix-up, no chunks touched',
        'Layout and dating are unaffected either way',
    ], space_after=9)
    footer(s, n)
    notes(s, 'If a field is an hourly mean, its stamp means "the hour ending at", which belongs '
             'in the attributes and the HOWTO caveat. The data go in the same hourly store '
             'regardless.')


def s_append(prs, n):
    s = new_slide(prs)
    title(s, 'Design change 1: append, never rewrite',
          'Stores open in mode="a", completeness is tracked per variable, and masks follow the grid location')
    card(s, 0.6, 1.85, 6.0, 4.9, 'mode="a" + per-variable complete', [
        'Create only the arrays that are absent or incomplete',
        'Each array: complete = true after verification',
        'Group complete = all listed variables complete',
        'Existing 9,501 Theta stores valid untouched',
        'A missing field no longer forces a full rewrite',
    ], space_after=9)
    card(s, 6.8, 1.85, 5.95, 4.9, 'Mask by staggering', [
        'i_g in dims → hFacW;  j_g → hFacS;  else hFacC',
        'Fixes oceTAUX/Y and SIuice/vice (were hFacC)',
        '--mask-dir mandatory: zeros are real for ice fields',
        'Streaming: read, mask, write, verify, one field at a time',
        '~3 GB per process, not ~20 GB',
    ], space_after=9)
    footer(s, n)
    notes(s, 'store_is_complete(url, vars) asks only about the requested variables; legacy rule: '
             'group complete and listed in attrs["variables"]. Retries wrap a single variable. '
             'Two processes appending to one group would race on the group document, hence the '
             'one-writer rule enforced by --workers.')


def s_parallel(prs, n):
    s = new_slide(prs)
    title(s, 'Design change 2: more requests in flight',
          'Raise zarr concurrency inside each process and run several processes over disjoint date windows')
    cards = [
        ('--concurrency 32', ['zarr async.concurrency (default 10)',
                              'whole-variable write + read-back',
                              'expect 85 → 30–45 s per field-hour']),
        ('--workers N', ['hours split into N contiguous blocks',
                         'one process per block, own connection pool',
                         'no store ever visited by two workers']),
        ('--verify', ['full: read everything back (default)',
                      'sample: one face per variable',
                      'none: trust the PUT']),
    ]
    for k, (head, items) in enumerate(cards):
        x = 0.6 + k * 4.15
        card(s, x, 1.85, 3.9, 3.6, head, items, head_size=22, space_after=8)
    text(s, 0.6, 5.65, 12.1, 1.1, [
        [('Rules: ', {'bold': True, 'color': CORAL}),
         ('one extraction command at a time; no repair or consolidate tool against the bucket '
          'while it runs; retries per variable', {})]], size=20)
    footer(s, n)
    notes(s, 'Threads of our own are not added: zarr\'s asyncio already provides the I/O '
             'concurrency and compression runs in its thread pool. Node: a PBS node '
             '(select=1:ncpus=8:model=bro) if compute nodes reach Nautilus, else the PFE with '
             'at most 4 workers. ~3 GB per process per field in flight.')


def s_perf(prs, n):
    s = new_slide(prs)
    title(s, 'Why parallel helps: the loop waits on Nautilus',
          'Measured 85 s per field-hour with a 30 s floor from the same node; CPU and Lustre are ~10 s (inferred)')
    cd = CategoryChartData()
    cd.categories = ['typical (1,028 stores)', 'later 24 stores', 'best burst (5 stores)',
                     'endpoint contention']
    cd.add_series('seconds per store', (74.6, 82.4, 30, 133))
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.5), Inches(1.75),
                            Inches(7.4), Inches(5.1), cd)
    ch = gf.chart
    ch.has_legend = False
    ch.has_title = False
    ch.font.size = Pt(20)
    ch.font.name = BODY
    plot = ch.plots[0]
    plot.gap_width = 60
    plot.has_data_labels = True
    plot.data_labels.font.size = Pt(20)
    plot.data_labels.number_format = '0'
    plot.data_labels.number_format_is_linked = False
    plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
    ser = plot.series[0]
    ser.format.fill.solid()
    ser.format.fill.fore_color.rgb = TEAL
    for idx, col in ((2, CORAL), (3, MUTED)):
        pt = ser.points[idx]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = col
    va, ca = ch.value_axis, ch.category_axis
    va.maximum_scale, va.minimum_scale, va.major_unit = 150, 0, 50
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = GRID
    va.tick_labels.font.size = Pt(20)
    ca.tick_labels.font.size = Pt(20)
    ca.has_major_gridlines = False
    va.has_title = True
    va.axis_title.text_frame.text = 's per store'
    r = va.axis_title.text_frame.paragraphs[0].runs[0]
    r.font.size, r.font.bold = Pt(20), False
    text(s, 8.2, 1.85, 4.6, 5.0, [
        [('11 MB/s ', {'bold': True}), ('on the wire: each byte PUT, then GET back', {})],
        [('≤ 10 ', {'bold': True}), ('of ~400 PUTs in flight; GETs never overlap', {})],
        [('~1 s ', {'bold': True}), ('per round trip: latency-bound', {})],
        [('120 MB/s ', {'bold': True}), ('reached by 6 workers from a workstation', {})],
    ], size=20, bullets=True, space_after=12)
    footer(s, n)
    notes(s, 'Measured from object timestamps (llc4320_v2.md logs 2026-09-25, 10-02, 10-09). '
             'The 30 s burst (2026-09-30) caps the whole cycle; the rest is S3 wait. The 133 s '
             'reading was while the face repair ran 6 workers against the endpoint. Inferred: '
             'Lustre read 1-3 s, numpy 2-3 s, zstd 2-4 s.')


def s_bench(prs, n):
    s = new_slide(prs)
    title(s, 'Benchmark before launch (about 40 minutes)',
          'Three timing runs into a throwaway prefix and one local run that measures per-field sizes')
    runs = [
        ('A', 'Baseline', '1 process, defaults, 6 hours', 'expect ~85 s/store'),
        ('B', 'Concurrency', 'ZARR_ASYNC__CONCURRENCY=40, 6 other hours', '35–45 s confirms latency-bound'),
        ('C', 'Four windows', '4 processes, disjoint 6-hour windows', 'near B/4 → use 6–8 workers'),
        ('D', 'Sizes', '1 hour, all 21 fields, local scratch', 'du per variable → storage total'),
    ]
    for k, (letter, head, what, reading) in enumerate(runs):
        x = 0.6 + (k % 2) * 6.2
        y = 1.85 + (k // 2) * 2.5
        box(s, x, y, 5.9, 2.25, ICE if k < 3 else BLUSH)
        dot(s, x + 0.3, y + 0.3, letter, TEAL if k < 3 else CORAL)
        text(s, x + 1.0, y + 0.25, 4.7, 0.55, [head], size=24, color=NAVY, font=HEAD)
        text(s, x + 1.0, y + 0.9, 4.7, 1.3, [what, [(reading, {'color': MUTED})]], size=20,
             space_after=4)
    footer(s, n)
    notes(s, 'wr_llc_v2_surface $OUT s3://llc4320-v2/BENCH --fields SST --no-grid --mask-dir '
             '$MASKDIR, different hours per run (Lustre caches re-reads). Also qsub -v LIMIT=1 to '
             'learn whether compute nodes reach Nautilus. Then aws s3 rm --recursive the BENCH '
             'prefix. Record A-D in llc4320_v2_other_fields.md Logs.')


def s_order(prs, n):
    s = new_slide(prs)
    title(s, 'Run order',
          'Close the SST gaps now; code; checks; benchmark; one pass; the last 189 hours when Dan opens them')
    steps = [
        ('Two outage hours', 'current code, SST only, minutes'),
        ('Code + tests', 'append mode, masks, workers, verify'),
        ('Pre-launch pastes', 'folder span, diag frequencies, PBS egress'),
        ('Storage decision', 'quota ~70–80 TB or trim the list'),
        ('Benchmark', 'A–D, pick N and verify mode'),
        ('One pass, all fields', '--fields all --workers N; audit after'),
        ('HOWTO + docs', 'tables, rotation example, coverage'),
        ('189 hours', 'same command, --start 2024-02-01'),
    ]
    for k, (head, body) in enumerate(steps):
        x = 0.6 + (k % 4) * 3.1
        y = 2.0 + (k // 4) * 2.5
        dot(s, x, y, k + 1, CORAL if k in (3, 5) else TEAL)
        text(s, x + 0.65, y - 0.02, 2.4, 0.9, [head], size=22, color=NAVY, font=HEAD)
        text(s, x, y + 0.75, 2.95, 1.5, [body], size=20, color=INK)
    footer(s, n)
    notes(s, 'Step 6 also extends the SST record if the tree now runs past 2024-02-08, since the '
             'pass writes all 21 variables including Theta for any new hours. Re-running the same '
             'command finishes whatever failed; it is idempotent per variable.')


def s_cost(prs, n):
    s = new_slide(prs)
    title(s, 'Storage and runtime',
          'Every estimate exceeds the 55 TB accepted; with ~6 workers the pass takes about two weeks')
    stat(s, 0.6, 1.9, 2.9, '55 TB', 'accepted so far', color=MUTED)
    stat(s, 3.6, 1.9, 3.1, '60–65 TB', 'likely total (21 vars)', color=CORAL)
    stat(s, 6.9, 1.9, 2.9, '97 TB', 'upper bound', color=CORAL)
    stat(s, 9.9, 1.9, 2.9, '×1.15', 'if ~11,000 hours', color=MUTED)
    text(s, 0.6, 3.35, 12.1, 0.6, [
        [('12 dense fields alone ≈ 55 TB (inferred from SST\'s 0.47 GB per field-hour); '
          'the sparse ice fields are 1–4 TB if they compress 10–50×, 37 TB if they do not', {})]],
         size=20, color=MUTED)
    table(s, 0.6, 4.2, 12.1, [
        ['Configuration (20 fields, 9,692 h)', 's per field-hour', 'Wall clock'],
        ['As the code stands, 1 process', '85 (measured)', '~190 days'],
        ['Concurrency 32, whole-variable writes, 1 process', '40 (inferred)', '~90 days'],
        ['Same, 6 workers', '40', '~15 days'],
        ['Same, 6 workers, --verify sample', '~30', '~11 days'],
    ], col_w=[6.6, 2.6, 2.9], row_h=0.5)
    footer(s, n)
    notes(s, 'Objects: ~82 M in total; a sparse field costs the same round trips as a dense one. '
             'If trimming, PhiBot and KPPhbl (4.6 TB each) have the least surface-ocean interest; '
             'the sparse fields are not worth dropping for storage. Benchmark run D measures the '
             'sparse compression before anything is written.')


def s_risks(prs, n):
    s = new_slide(prs, dark=True)
    title(s, 'Risks and what Xavier decides next',
          'Two decisions and three pastes stand between the design and the launch', dark=True)
    items = [
        ('Storage over 55 TB', 'confirm ~70–80 TB with Nautilus, or trim'),
        ('95 folders', 'paste span + dry-run count; proof redone'),
        ('Snapshot vs. mean', 'grep frequencies, or ask Dan'),
        ('Endpoint throttling', 'retries per variable; lower N if failures climb'),
        ('Two writers', '--workers partitions; one command at a time'),
        ('Dan', 'open run_sep32026; SIuice/SIvice points; R_* files'),
    ]
    for k, (head, body) in enumerate(items):
        x = 0.6 + (k % 2) * 6.25
        y = 1.9 + (k // 2) * 1.65
        dot(s, x, y + 0.05, k + 1, CORAL if k < 3 else TEAL)
        text(s, x + 0.75, y, 5.3, 0.55, [head], size=24, color=WHITE, font=HEAD)
        text(s, x + 0.75, y + 0.62, 5.3, 0.9, [body], size=20, color=ICE)
    footer(s, n, dark=True)
    notes(s, 'Also: compute-node egress to Nautilus unknown (qsub -v LIMIT=1); mixed variable '
             'sets during the run (open_range intersection); the exact-zero artefact repeats in '
             'every hFacC-masked field; field list settled by the inventory but still resolved at '
             'run time as a guard.')


def main():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    prs.core_properties.title = 'LLC4320 v2: the other surface fields -- design'
    prs.core_properties.author = 'J. Xavier Prochaska'
    s_title(prs)
    for n, fn in enumerate([s_goal, s_source, s_pipeline, s_faces, s_grid, s_access, s_fields,
                            s_inventory, s_sampling, s_append, s_parallel, s_perf, s_bench,
                            s_order, s_cost, s_risks], start=2):
        fn(prs, n)
    prs.save(OUT)
    print(OUT)


if __name__ == '__main__':
    main()
