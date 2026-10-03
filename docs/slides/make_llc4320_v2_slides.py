""" Build docs/slides/llc4320_v2.pptx (status of the LLC4320 v2 -> Nautilus work).

Run llc4320_v2_figures.py first (it writes figs/*.png and figs/progress.csv),
then:

    python docs/slides/make_llc4320_v2_slides.py

Needs python-pptx and Pillow.  The deck is meant to be edited by hand in
LibreOffice Impress afterwards; re-running this overwrites those edits.

Every font on the slides is at least MIN_PT (20 pt); `_run` raises if a
smaller size slips in.  Speaker notes carry the detail that no longer fits.
"""

import csv
import datetime as dt
import os

from PIL import Image
from pptx import Presentation
from pptx.chart.data import XyChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
FIGS = os.path.join(HERE, 'figs')
OUT = os.path.join(HERE, 'llc4320_v2.pptx')

# Palette: deep-ocean navy dominates, warm SST coral as the single accent.
NAVY = RGBColor(0x0B, 0x2A, 0x3F)
TEAL = RGBColor(0x1F, 0x6F, 0x9F)
ICE = RGBColor(0xDC, 0xEB, 0xF2)
CORAL = RGBColor(0xE4, 0x57, 0x2E)
INK = RGBColor(0x1C, 0x26, 0x30)
MUTED = RGBColor(0x5B, 0x6B, 0x78)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
HEAD, BODY, MONO = 'Cambria', 'Calibri', 'Courier New'

MIN_PT = 20          # smallest font allowed anywhere on a slide
BODY_PT = 20
TITLE_PT = 36

W, H = 13.333, 7.5
STATUS_DATE = '2026-10-02'
TOTAL_HOURS = 9503


def _run(p, text, size, color=INK, bold=False, italic=False, font=BODY):
    if size < MIN_PT:
        raise ValueError(f'{size} pt is below the {MIN_PT} pt minimum: {text!r}')
    r = p.add_run()
    r.text = text
    f = r.font
    f.size, f.bold, f.italic, f.name = Pt(size), bold, italic, font
    f.color.rgb = color
    return r


def text(slide, x, y, w, h, paras, size=BODY_PT, color=INK, font=BODY, anchor=MSO_ANCHOR.TOP,
         align=PP_ALIGN.LEFT, space_after=8, bullets=False):
    """Text box; *paras* is a list of str or list of (text, {opts}) runs."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for n, para in enumerate(paras):
        p = tf.paragraphs[0] if n == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        runs = [(para, {})] if isinstance(para, str) else para
        if bullets:
            runs = [('•  ', {'color': CORAL, 'bold': True})] + list(runs)
        for t, o in runs:
            _run(p, t, o.get('size', size), o.get('color', color), o.get('bold', False),
                 o.get('italic', False), o.get('font', font))
    return tb


def box(slide, x, y, w, h, fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    s.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = 0.08
    return s


def numbered_dot(slide, x, y, n, fill, d=0.5):
    dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    dot.fill.solid()
    dot.fill.fore_color.rgb = fill
    dot.line.fill.background()
    tf = dot.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    _run(p, str(n), 20, WHITE, bold=True)
    return dot


def background(slide, color):
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = color


def title(slide, t, sub=None, dark=False, size=TITLE_PT):
    text(slide, 0.6, 0.4, W - 1.2, 0.8, [t], size=size, color=WHITE if dark else NAVY, font=HEAD)
    if sub:
        text(slide, 0.6, 1.2, W - 1.2, 0.5, [sub], size=20, color=ICE if dark else MUTED)


def footer(slide, n, dark=False):
    """Slide number only: a 20 pt running footer would crowd the content."""
    text(slide, W - 1.4, H - 0.55, 0.9, 0.35, [str(n)], size=20, color=ICE if dark else MUTED,
         align=PP_ALIGN.RIGHT)


def picture(slide, path, x, y, w, h):
    """Insert *path* scaled to fit the (x, y, w, h) box, centred."""
    iw, ih = Image.open(path).size
    s = min(w / iw, h / ih)
    pw, ph = iw * s, ih * s
    return slide.shapes.add_picture(path, Inches(x + (w - pw) / 2), Inches(y + (h - ph) / 2),
                                    Inches(pw), Inches(ph))


def stat(slide, x, y, w, value, label, color=NAVY, label_color=MUTED, size=40):
    text(slide, x, y, w, 0.75, [value], size=size, color=color, font=HEAD, space_after=0)
    text(slide, x, y + 0.75, w, 0.45, [label], size=20, color=label_color, space_after=0)


def notes(slide, t):
    slide.notes_slide.notes_text_frame.text = t


def read_progress():
    path = os.path.join(FIGS, 'progress.csv')
    with open(path, newline='', encoding='utf-8') as fh:
        rows = list(csv.DictReader(fh))
    for r in rows:
        r['written_utc'] = dt.datetime.strptime(r['written_utc'], '%Y-%m-%d %H:%M')
        r['n_stores'] = int(r['n_stores'])
    return rows


# ------------------------------------------------------------------ slides

def s_title(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    background(s, NAVY)
    picture(s, os.path.join(FIGS, 'llc4320_v2_sst_global.png'), 6.9, 1.3, 6.0, 4.6)
    text(s, 0.6, 1.8, 6.0, 1.8, ['LLC4320 v2', 'surface fields on Nautilus'], size=40,
         color=WHITE, font=HEAD, space_after=0)
    text(s, 0.6, 3.8, 5.9, 1.2,
         ['Raw MITgcm output on NASA Pleiades → hourly Zarr stores in s3://llc4320-v2'],
         size=22, color=ICE)
    text(s, 0.6, 5.5, 6, 0.8, [f'J. X. Prochaska  ·  {STATUS_DATE}'], size=20, color=ICE)
    notes(s, 'Summary of the LLC4320 v2 extraction work: what the run is, how the pipeline '
             'works, what the data look like, and where the full SST run stands.')


def s_run(prs, n):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    background(s, WHITE)
    title(s, 'The v2 simulation', 'A new LLC4320 run by Dan Whitt (NASA Ames), not the v1 data in s3://dbof')
    text(s, 0.6, 2.0, 6.3, 4.8, [
        [('Global LLC grid, ', {'bold': True}), ('1/48°', {})],
        [('Same horizontal grid as v1; ', {'bold': True}), ('new bathymetry, 173 levels', {})],
        [('Hourly snapshots ', {'bold': True}), ('from 2023-01-01; the run is still going', {})],
        [('83+ folders ', {'bold': True}), ('on Pleiades; model ΔT 5–20 s', {})],
        [('Goal: ', {'bold': True}), ('all surface fields on Nautilus, SST first', {})],
    ], size=22, bullets=True, space_after=14)
    box(s, 7.3, 1.95, 5.4, 4.85, ICE).name = 'stats card'
    stat(s, 7.7, 2.15, 2.4, '9,503', 'hourly fields')
    stat(s, 10.2, 2.15, 2.4, '83', 'folders')
    stat(s, 7.7, 3.7, 2.4, '13×4320²', 'cells per field', size=32)
    stat(s, 10.2, 3.7, 2.4, '58 %', 'ocean')
    stat(s, 7.7, 5.25, 2.4, '970 MB', 'raw per hour', size=32)
    stat(s, 10.2, 5.25, 2.4, '474 MB', 'Zarr per hour', size=32)
    footer(s, n)
    notes(s, '13 faces x 4320 x 4320 = 243 M cells; 58 % ocean. 970 MB is one raw SST file; '
             '474 MB is one compressed hourly Zarr store on S3. Data on disk run to 2024-02-08.')


def s_formats(prs, n):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    background(s, WHITE)
    title(s, 'Two raw formats; the easy one covers the surface')
    cols = [
        ('Surface: .data', TEAL, ICE, [
            'SST, SSS, SSU, SSV, Eta, fluxes, sea ice',
            'Uncompressed, big-endian float32',
            'In all 83 folders',
            [('Used now', {'bold': True, 'color': TEAL})],
        ]),
        ('3-D: .shrunk', MUTED, RGBColor(0xEE, 0xF1, 0xF3), [
            'Theta, Salt, U, V, W on 173 levels',
            'Wet points only + bitmasks',
            'Jan–Jun 2023 moved to NAS export2',
            [('Decoder ready for the depth phase', {'bold': True, 'color': MUTED})],
        ]),
    ]
    for k, (head, col, fill, items) in enumerate(cols):
        x = 0.6 + k * 6.25
        box(s, x, 1.6, 5.9, 5.0, fill)
        text(s, x + 0.4, 1.9, 5.2, 0.6, [head], size=28, color=col, font=HEAD)
        text(s, x + 0.4, 2.8, 5.2, 3.7, items, size=22, bullets=True, space_after=16)
    footer(s, n)
    notes(s, 'Surface files: SST/SSS/SSU/SSV.<iter>.data, MITgcm compact 4320 x 56160 real*4, '
             'identical to k=0 of the 3-D fields (Dan Whitt). 3-D files: <field>.<iter>.shrunk '
             'with hFacC/S/W.bits masks; only ~33 folders on POSIX. The decoder is a NumPy port '
             'of Kaitlin Zhang\'s llc_shrunk_mex.c, round-trip tested.')


def s_pipeline(prs, n):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    background(s, WHITE)
    title(s, 'The pipeline: wrangler.ogcm.llc_v2',
          'One process on a Pleiades login node, streaming to Nautilus')
    steps = [
        ('Discover', 'date = start + (n+1) h'),
        ('Read', 'SST.<iter>.data → 13 faces'),
        ('Mask', 'hFacC.bits → NaN on land'),
        ('Write', 'Zarr v3, verified'),
        ('Publish', 'one store per hour on S3'),
    ]
    bw, gap, y = 2.2, 0.32, 2.0
    for k, (head, body) in enumerate(steps):
        x = 0.6 + k * (bw + gap)
        dark = k in (0, 4)
        box(s, x, y, bw, 2.3, NAVY if dark else ICE)
        text(s, x + 0.2, y + 0.2, bw - 0.4, 0.5, [head], size=24, font=HEAD,
             color=WHITE if dark else NAVY)
        text(s, x + 0.2, y + 0.85, bw - 0.4, 1.4, [body], size=20,
             color=ICE if dark else INK)
        if k < len(steps) - 1:
            a = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x + bw + 0.04), Inches(y + 1.0),
                                   Inches(gap - 0.08), Inches(0.35))
            a.fill.solid()
            a.fill.fore_color.rgb = CORAL
            a.line.fill.background()
    text(s, 0.6, 4.65, 6, 0.5, ['Command-line tools'], size=24, color=NAVY, font=HEAD)
    text(s, 0.6, 5.25, 6.4, 1.6, [
        [('wr_llc_v2_inventory', {'font': MONO, 'bold': True})],
        [('wr_llc_v2_surface', {'font': MONO, 'bold': True})],
        [('wr_llc_v2_inspect', {'font': MONO, 'bold': True})],
    ], size=20, space_after=4)
    text(s, 7.2, 4.65, 5.5, 0.5, ['Built to be restarted'], size=24, color=NAVY, font=HEAD)
    text(s, 7.2, 5.25, 5.5, 1.6, [
        'Finished stores are skipped',
        'Unreadable folders are skipped',
    ], size=20, bullets=True, space_after=6)
    footer(s, n)
    notes(s, 'Discover: list segment folders, date each file as folder start + (n+1) h, strict '
             'file-count check. Write: chunks (1, 720, 720), read back and verified; '
             'complete=True is set last, so partial stores are rewritten. Stores are '
             's3://llc4320-v2/SURFACE/YYYYMMDDTHH.zarr; grid.zarr holds 23 static fields. '
             'Tools: inventory = what is in each folder; surface = extract (--fields, --dry-run, '
             '--limit); inspect = per-face stats, mask check, PNG.')


def s_lessons(prs, n):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    background(s, WHITE)
    title(s, 'Problems found and fixed along the way')
    cards = [
        ('Dates one hour early', 'Fixed: 9,503 hours, no gaps or repeats'),
        ('Unreadable folders', 'The live run and a locked data/ dir are now skipped'),
        ('Missing 3-D output', 'Jan–Jun 2023 .shrunk moved to export2'),
        ('3,477 wet cells = 0', '0.0025 % of ocean; mask vs. bathymetry'),
    ]
    for k, (head, body) in enumerate(cards):
        x = 0.6 + (k % 2) * 6.25
        y = 1.5 + (k // 2) * 2.7
        box(s, x, y, 5.9, 2.45, ICE)
        numbered_dot(s, x + 0.3, y + 0.32, k + 1, CORAL if k == 3 else TEAL)
        text(s, x + 1.0, y + 0.3, 4.7, 0.6, [head], size=24, color=NAVY, font=HEAD)
        text(s, x + 1.0, y + 1.0, 4.65, 1.35, [body], size=20)
    footer(s, n)
    notes(s, '1. ExtractFields.m dated files as start + n h; Dan: the first file is one hour '
             'after the folder start. 9,503 files now fill 2023-01-01 01:00 to 2024-01-31 23:00. '
             '2. The newest folder is a symlink into the live run; one folder has an unreadable '
             'data/ directory. 3. Theta.*.shrunk was missing from early folders: it moved to the '
             'export2 object store (mc + NAS access). 4. Cells the bitmask calls wet but the run '
             'wrote as 0 (bathymetry mismatch or ice-shelf cavities); left as 0.0.')


def s_global(prs, n):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    background(s, WHITE)
    title(s, 'On Nautilus: one hour of global SST')
    picture(s, os.path.join(FIGS, 'llc4320_v2_sst_global.png'), 0.5, 1.35, 8.9, 5.6)
    text(s, 9.7, 1.8, 3.1, 4.8, [
        '2023-07-01 12:00 UTC',
        'Every 4th point, binned to 0.25°',
        '−2.1 to 32.4 °C',
    ], size=22, bullets=True, space_after=20)
    footer(s, n)
    notes(s, 'Read from s3://llc4320-v2/SURFACE/20230701T12.zarr by docs/slides/'
             'llc4320_v2_figures.py. Range and mean (11.85 C) are from the first S3 store; '
             'nothing outside the plausible range.')


def s_zoom(prs, n):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    background(s, WHITE)
    title(s, 'At full resolution: the Gulf Stream')
    picture(s, os.path.join(FIGS, 'llc4320_v2_sst_gulfstream.png'), 0.5, 1.35, 9.0, 5.6)
    text(s, 9.8, 1.8, 3.0, 4.8, [
        'Every cell, ~2 km',
        'Same hour',
        'Fronts, filaments and eddies',
    ], size=22, bullets=True, space_after=20)
    footer(s, n)
    notes(s, '960 x 804 cells of face 10, no subsampling. Small-scale structure like this is '
             'the reason to keep the native grid.')


def s_layout(prs, n):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    background(s, WHITE)
    title(s, 'Found this week: faces 7–12 were scrambled', size=34)
    picture(s, os.path.join(FIGS, 'llc4320_v2_face_layout.png'), 0.5, 1.35, 7.6, 5.6)
    text(s, 8.5, 1.6, 4.3, 5.2, [
        [('Cause: ', {'bold': True}), ('wrong reshape of the rotated faces', {})],
        [('Nothing lost: ', {'bold': True}), ('values are only permuted', {})],
        [('Fixed: ', {'bold': True}), ('matches v1\'s grid exactly', {})],
        [('To do: ', {'bold': True}), ('repair 5,493 stores', {})],
    ], size=22, bullets=True, space_after=18)
    footer(s, n)
    notes(s, 'Top row: faces 10-12 of the 2023-07-01 12 UTC store as written. Bottom row: '
             'after llc_v2.compact_to_faces. The compact file was reshaped straight to '
             '(13, 4320, 4320); that is right for faces 0-6, but faces 7-12 are stored as '
             '(4320, 12960) blocks. The fix reproduces XC/YC of s3://dbof grid.zarr with max '
             'difference 0. Fixed stores carry face_layout = "llc_faces". The global map was '
             'unaffected because it bins by XC/YC.')


def s_series(prs, n):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    background(s, WHITE)
    title(s, 'Through time: the dates line up',
          'Left: hourly for one week.  Right: daily at 12 UTC, 2023-01 to 08')
    picture(s, os.path.join(FIGS, 'llc4320_v2_sst_series.png'), 0.5, 1.85, 12.3, 5.1)
    footer(s, n)
    notes(s, 'Single grid cells read from the hourly stores. A smooth diurnal cycle and a '
             'seasonal cycle with no jumps at folder boundaries both support the dating rule.')


def s_progress(prs, n):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    background(s, WHITE)
    rows = read_progress()
    done = rows[-1]['n_stores']
    last_t = rows[-1]['written_utc']
    title(s, 'The full SST run: past halfway, now stalled')
    cd = XyChartData()
    ser = cd.add_series('stores written')
    t0 = rows[0]['written_utc']
    for r in rows:
        ser.add_data_point(round((r['written_utc'] - t0).total_seconds() / 86400, 3), r['n_stores'])
    gf = s.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER_LINES_NO_MARKERS, Inches(0.5), Inches(1.4),
                            Inches(7.7), Inches(5.5), cd)
    ch = gf.chart
    ch.has_legend = False
    ch.font.size = Pt(20)
    ch.font.name = BODY
    ch.series[0].format.line.color.rgb = TEAL
    ch.series[0].format.line.width = Pt(3)
    ch.series[0].smooth = False
    va, ca = ch.value_axis, ch.category_axis
    va.maximum_scale, va.minimum_scale, va.major_unit = 10000, 0, 2500
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = RGBColor(0xDD, 0xE3, 0xE8)
    va.tick_labels.number_format = '#,##0'
    va.tick_labels.number_format_is_linked = False
    va.has_title = True
    va.axis_title.text_frame.text = 'stores written'
    ca.minimum_scale, ca.maximum_scale, ca.major_unit = 0, 8, 2
    ca.has_major_gridlines = False
    ca.has_title = True
    ca.axis_title.text_frame.text = 'days since 2026-09-24'
    for ax in (va, ca):
        r = ax.axis_title.text_frame.paragraphs[0].runs[0]
        r.font.size = Pt(20)
        r.font.bold = False

    x = 8.7
    stat(s, x, 1.5, 4.1, f'{done:,}', f'of {TOTAL_HOURS:,} stores ({100 * done / TOTAL_HOURS:.0f} %)')
    stat(s, x, 2.85, 4.1, 'Jan 1 → Aug 17', 'model dates, no gaps', size=32)
    stat(s, x, 4.2, 4.1, '4.5 TB', 'projected total', size=32)
    box(s, x - 0.1, 5.55, 4.2, 1.3, RGBColor(0xFB, 0xE4, 0xDC)).name = 'stall alert'
    text(s, x + 0.15, 5.7, 3.8, 1.1, [
        [('Nothing new since', {'bold': True})],
        [(f'{last_t:%b %d %H:%M} UTC', {'bold': True, 'color': CORAL})],
    ], size=22, space_after=2)
    footer(s, n)
    notes(s, f'Read directly from S3 on {STATUS_DATE}. Launched 2026-09-24 with --fields SST. '
             '~75 s and 474 MB per store while running. It stopped on 09-29 14:03 UTC after '
             '20230817T16; 5 more stores were written 09-30 23:01-23:04 UTC, then nothing.')


def s_next(prs, n):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    background(s, NAVY)
    title(s, 'Next steps', dark=True)
    items = [
        ('Restart with the fix', 'git pull on Pleiades first'),
        ('Repair old stores', 'rewrite faces 7–12 on S3'),
        ('Bucket quota', '4.5 TB for SST; ~18 TB for 4 fields'),
        ('SSS / SSU / SSV?', 'each later pass is ~8 days'),
        ('Depth phase', 'needs NAS export2 access'),
        ('Vectors', 'rotate SSU/SSV to east/north'),
    ]
    for k, (head, body) in enumerate(items):
        x = 0.6 + (k % 2) * 6.25
        y = 1.55 + (k // 2) * 1.8
        numbered_dot(s, x, y + 0.05, k + 1, CORAL if k < 2 else TEAL)
        text(s, x + 0.75, y, 5.2, 0.55, [head], size=26, color=WHITE, font=HEAD)
        text(s, x + 0.75, y + 0.65, 5.2, 0.9, [body], size=20, color=ICE)
    footer(s, n, dark=True)
    notes(s, 'Repair option: rewrite faces 7-12 of each existing store in place from a '
             'workstation (~2.6 TB of S3 traffic, no Pleiades reads), or re-extract them '
             '(~4.7 more days on Pleiades). Adding fields later rewrites every store. '
             'Vectors need AngleCS/AngleSN from grid.zarr.')


def main():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    prs.core_properties.title = 'LLC4320 v2 surface fields on Nautilus'
    prs.core_properties.author = 'J. Xavier Prochaska'
    s_title(prs)
    for n, fn in enumerate([s_run, s_formats, s_pipeline, s_lessons, s_global, s_zoom,
                            s_layout, s_series, s_progress, s_next], start=2):
        fn(prs, n)
    prs.save(OUT)
    print(OUT)


if __name__ == '__main__':
    main()
